"""Stage C Task 29 — CompPhish mapping schema + final-holdout identity index.

Reads only identity fields from the sealed CompPhish v3 mapping workbook,
reconciles them one-to-one with the sealed HTML archive, hashes URLs instead of
persisting them, and hashes each HTML capture for later contamination audits.
No production feature extraction, model scoring, metrics, threshold changes,
or deployment occur here.
"""
from __future__ import annotations

from collections import Counter
from decimal import Decimal, InvalidOperation
import hashlib
import json
import os
from pathlib import Path, PurePosixPath
import re
import subprocess
from typing import Any, Mapping
from urllib.parse import urlsplit, urlunsplit
import zipfile

try:
    from openpyxl import load_workbook
except ImportError as exc:  # fail clearly on the user's runtime if dependency is absent
    raise RuntimeError('Task 29 requires openpyxl; install it before running this gate') from exc

INDEX_SCHEMA = 'stage-c-final-holdout-identity-index-1'
REPORT_SCHEMA = 'stage-c-final-holdout-mapping-schema-report-1'

EXPECTED_TASK28_SEAL_SHA256 = 'bfd50f0ded9819f8212c285e4a9821834b3c67beff79d56cdc8c43d1c602199a'
EXPECTED_TASK28_ARTIFACT_SET_SHA256 = '15bb4b7c1ad8536b68b88a64fbd2ea8bc68de49805a25c53df57cfbc0919724f'
EXPECTED_HTML_SHA256 = '12440e4f911fabf4ec2c712ae014cb43e638e9a6adc6c7c6a506a8a7df24fcf7'
EXPECTED_HTML_MEMBER_INVENTORY_SHA256 = '16cb22c754e685ba7310a5b021d82f9e11dd4cc3341e5841b1998fba38edcdc0'
EXPECTED_MAPPING_SHA256 = 'e6dda6a21d925f8aba36090f295b08e9808a5dfe8dd3276bf8cd82022855f14a'

EXPECTED_CANDIDATE_ID = 'compphish-v3-2026'
EXPECTED_SOURCE_ID = 'mendeley-fmbs4kp9wz-v3'
EXPECTED_DOI = '10.17632/fmbs4kp9wz.3'
EXPECTED_HTML_ARCHIVE_NAME = 'All_HTML.zip'
EXPECTED_MAPPING_NAME = 'Mapping_File.xlsx'
EXPECTED_SAMPLE_COUNT = 15358
EXPECTED_CLASS_COUNTS = {0: 8154, 1: 7204}

HEADER_ALIASES = {
    'serial': {'serial_number','serial_no','serial','serialnumber','sr_no','sr_number','s_no','sno','id','index'},
    'url': {'url','urls','website_url','web_url','website','link'},
    'label': {'label','labels','class','target','result','status'},
    'html': {'html','html_file','html_filename','html_source','html_source_code','html_code_file','filename','file_name'},
}

class StageCFinalHoldoutIdentityIndexError(ValueError):
    pass

def canonical_bytes(value: Any) -> bytes:
    return json.dumps(value, sort_keys=True, separators=(',', ':'), ensure_ascii=False).encode('utf-8')

def canonical_hash(value: Any) -> str:
    return hashlib.sha256(canonical_bytes(value)).hexdigest()

def hash_without(value: Mapping[str, Any], field: str) -> str:
    copy = dict(value); copy.pop(field, None); return canonical_hash(copy)

def sha256_file(path: Path, chunk_size: int = 8 * 1024 * 1024) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(chunk_size), b''): h.update(chunk)
    return h.hexdigest()

def sha256_zip_member(archive: zipfile.ZipFile, info: zipfile.ZipInfo) -> str:
    h = hashlib.sha256()
    with archive.open(info, 'r') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''): h.update(chunk)
    return h.hexdigest()

def load_json(path: Path) -> dict[str, Any]:
    if not path.is_file(): raise StageCFinalHoldoutIdentityIndexError(f'required JSON file not found: {path}')
    value = json.loads(path.read_text(encoding='utf-8'))
    if not isinstance(value, dict): raise StageCFinalHoldoutIdentityIndexError(f'JSON root must be object: {path}')
    return value

def frozen_write_json(path: Path, value: Any) -> str:
    payload = (json.dumps(value, indent=2, sort_keys=True, ensure_ascii=False) + '\n').encode('utf-8')
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        if path.read_bytes() != payload: raise StageCFinalHoldoutIdentityIndexError(f'refusing to replace non-identical frozen Task-29 output: {path}')
        return 'EXISTING_MATCH'
    tmp = path.with_name(f'.{path.name}.tmp-{os.getpid()}')
    try:
        with tmp.open('xb') as f:
            f.write(payload); f.flush(); os.fsync(f.fileno())
        os.replace(tmp, path)
    finally: tmp.unlink(missing_ok=True)
    return 'CREATED'

def validate_task28_seal(seal: Mapping[str, Any]) -> None:
    expected = {
        'schema_version':'stage-c-final-holdout-local-integrity-seal-1','status':'PASS','stage':'C','role':'FINAL_HOLDOUT',
        'protocol_id':'low-fpr-generalization-v1','research_only':True,'deployment_authorized':False,
        'candidate_id':EXPECTED_CANDIDATE_ID,'source_id':EXPECTED_SOURCE_ID,'doi':EXPECTED_DOI,
        'expected_sample_count':EXPECTED_SAMPLE_COUNT,'expected_class_counts':{'legitimate':8154,'phishing':7204},
        'local_artifact_set_sha256':EXPECTED_TASK28_ARTIFACT_SET_SHA256,'local_integrity_seal_complete':True,
        'html_archive_integrity_passed':True,'mapping_container_integrity_passed':True,'mapping_schema_inspected':False,
        'mapping_labels_accessed':False,'html_capture_contents_accessed':False,'final_holdout_bytes_accessed_for_integrity':True,
        'final_holdout_touch_scope':'BYTE_AND_CONTAINER_INTEGRITY_ONLY','final_holdout_feature_extraction_authorized':False,
        'final_holdout_model_scoring_authorized':False,'final_holdout_metrics_authorized':False,
        'local_integrity_seal_sha256':EXPECTED_TASK28_SEAL_SHA256,
        'next_gate':'INSPECT_COMPPHISH_MAPPING_SCHEMA_AND_BUILD_FINAL_HOLDOUT_IDENTITY_INDEX_WITHOUT_MODEL_SCORING',
    }
    for key, value in expected.items():
        if seal.get(key) != value: raise StageCFinalHoldoutIdentityIndexError(f'Task-28 seal guard mismatch: {key}')
    if hash_without(seal, 'local_integrity_seal_sha256') != EXPECTED_TASK28_SEAL_SHA256:
        raise StageCFinalHoldoutIdentityIndexError('Task-28 canonical seal hash mismatch')
    artifacts = seal.get('local_artifacts')
    if not isinstance(artifacts, Mapping): raise StageCFinalHoldoutIdentityIndexError('Task-28 local_artifacts missing')
    html, mapping = artifacts.get('html_archive'), artifacts.get('mapping_workbook')
    if not isinstance(html, Mapping) or not isinstance(mapping, Mapping): raise StageCFinalHoldoutIdentityIndexError('Task-28 artifact identities missing')
    checks = [
        (html.get('filename')==EXPECTED_HTML_ARCHIVE_NAME,'HTML filename'),(html.get('sha256')==EXPECTED_HTML_SHA256,'HTML SHA-256'),
        (html.get('member_inventory_sha256')==EXPECTED_HTML_MEMBER_INVENTORY_SHA256,'HTML member inventory'),
        (html.get('txt_capture_count')==EXPECTED_SAMPLE_COUNT,'HTML capture count'),
        (mapping.get('filename')==EXPECTED_MAPPING_NAME,'mapping filename'),(mapping.get('sha256')==EXPECTED_MAPPING_SHA256,'mapping SHA-256')]
    for ok, name in checks:
        if not ok: raise StageCFinalHoldoutIdentityIndexError(f'Task-28 artifact identity changed: {name}')

def _normalize_header(value: str) -> str:
    value = re.sub(r'[\s\-/\\]+','_',value.strip().casefold())
    value = re.sub(r'[^a-z0-9_]+','',value)
    return re.sub(r'_+','_',value).strip('_')

def _resolve_columns(headers: list[str]) -> dict[str,int]:
    normalized = [_normalize_header(x) for x in headers]; resolved = {}
    for semantic, aliases in HEADER_ALIASES.items():
        matches = [i for i,n in enumerate(normalized) if n in aliases]
        if semantic == 'html':
            if len(matches)>1: raise StageCFinalHoldoutIdentityIndexError(f'ambiguous optional HTML column: {matches}')
            if matches: resolved[semantic]=matches[0]
        else:
            if len(matches)!=1: raise StageCFinalHoldoutIdentityIndexError(f'expected exactly one {semantic} column; headers={headers!r}, matches={matches}')
            resolved[semantic]=matches[0]
    return resolved

def _normalize_serial(value: Any) -> tuple[str,int|None]:
    raw = str(value).strip()
    if not raw: raise StageCFinalHoldoutIdentityIndexError('empty serial number')
    try: dec = Decimal(raw.replace(',',''))
    except InvalidOperation: return raw.casefold(), None
    if dec != dec.to_integral_value(): raise StageCFinalHoldoutIdentityIndexError(f'non-integral serial number: {raw!r}')
    integer = int(dec)
    if integer < 0: raise StageCFinalHoldoutIdentityIndexError(f'negative serial number: {raw!r}')
    return str(integer), integer

def _parse_label(value: Any) -> int:
    raw = str(value).strip().casefold()
    aliases = {'0':0,'0.0':0,'legitimate':0,'legit':0,'benign':0,'safe':0,'1':1,'1.0':1,'phishing':1,'phish':1,'malicious':1}
    if raw not in aliases: raise StageCFinalHoldoutIdentityIndexError(f'unsupported mapping label: {value!r}')
    return aliases[raw]

def _normalize_url(value: Any) -> tuple[str,str]:
    raw = str(value).strip()
    if not raw: raise StageCFinalHoldoutIdentityIndexError('empty URL')
    probe = raw if '://' in raw else f'http://{raw}'
    try: parts = urlsplit(probe); host = (parts.hostname or '').casefold().rstrip('.'); port = parts.port
    except ValueError as exc: raise StageCFinalHoldoutIdentityIndexError(f'invalid URL: {raw!r}') from exc
    if not host: raise StageCFinalHoldoutIdentityIndexError(f'URL has no hostname: {raw!r}')
    scheme = parts.scheme.casefold() or 'http'; netloc = host
    if port is not None and not ((scheme=='http' and port==80) or (scheme=='https' and port==443)): netloc=f'{host}:{port}'
    return urlunsplit((scheme,netloc,parts.path or '/',parts.query,'')), host

def _safe_member_name(name: str) -> bool:
    if not name or '\\' in name: return False
    p = PurePosixPath(name)
    return not p.is_absolute() and not any(x in {'','.','..'} for x in p.parts) and ':' not in p.parts[0]

def _sample_id(serial_key: str, serial_int: int|None) -> str:
    if serial_int is not None: return f'{EXPECTED_SOURCE_ID}:{serial_int:05d}'
    return f'{EXPECTED_SOURCE_ID}:{hashlib.sha256(serial_key.encode()).hexdigest()[:16]}'

def _git_provenance(repo_root: Path) -> dict[str,str]:
    paths=['ml/data/stage_c_final_holdout_identity_index.py','ml/data/build_stage_c_final_holdout_identity_index.py','ml/data/stage_c_final_holdout_local_seal.py']
    try:
        head=subprocess.run(['git','rev-parse','HEAD'],cwd=repo_root,capture_output=True,text=True,encoding='utf-8',check=True).stdout.strip()
        if len(head)!=40: raise StageCFinalHoldoutIdentityIndexError('invalid Git HEAD identity')
        for p in paths: subprocess.run(['git','cat-file','-e',f'HEAD:{p}'],cwd=repo_root,capture_output=True,check=True)
        dirty=subprocess.run(['git','diff','--quiet','HEAD','--',*paths],cwd=repo_root).returncode
    except (OSError,subprocess.CalledProcessError) as exc: raise StageCFinalHoldoutIdentityIndexError('Task-29 and bound files must be committed before identity-index build') from exc
    if dirty!=0: raise StageCFinalHoldoutIdentityIndexError('Task-29/bound files differ from committed HEAD')
    return {'git_head':head,'task29_module_sha256':sha256_file(repo_root/paths[0]),'task29_cli_sha256':sha256_file(repo_root/paths[1]),'task28_module_sha256':sha256_file(repo_root/paths[2])}

def _read_mapping(path: Path) -> tuple[list[str],list[list[Any]],str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(wb.sheetnames)!=1: raise StageCFinalHoldoutIdentityIndexError(f'expected exactly one worksheet, found {len(wb.sheetnames)}')
        ws=wb[wb.sheetnames[0]]; values=list(ws.iter_rows(values_only=True))
    finally: wb.close()
    values=[list(r) for r in values if any(v is not None and str(v).strip() for v in r)]
    if not values: raise StageCFinalHoldoutIdentityIndexError('mapping worksheet contains no data')
    headers=[str(x).strip() if x is not None else '' for x in values[0]]
    width=len(headers); rows=[]
    for r in values[1:]:
        r=list(r[:width])+[None]*max(0,width-len(r))
        if any(v is not None and str(v).strip() for v in r): rows.append(r)
    return headers, rows, ws.title

def build_final_holdout_identity_index(*, repo_root:Path, task28_seal_path:Path, html_archive_path:Path, mapping_workbook_path:Path) -> tuple[dict[str,Any],dict[str,Any]]:
    seal=load_json(task28_seal_path); validate_task28_seal(seal)
    if html_archive_path.name!=EXPECTED_HTML_ARCHIVE_NAME or not html_archive_path.is_file(): raise StageCFinalHoldoutIdentityIndexError('sealed HTML archive missing or misnamed')
    if mapping_workbook_path.name!=EXPECTED_MAPPING_NAME or not mapping_workbook_path.is_file(): raise StageCFinalHoldoutIdentityIndexError('sealed mapping workbook missing or misnamed')
    if sha256_file(html_archive_path)!=EXPECTED_HTML_SHA256: raise StageCFinalHoldoutIdentityIndexError('HTML archive SHA-256 differs from Task-28 seal')
    if sha256_file(mapping_workbook_path)!=EXPECTED_MAPPING_SHA256: raise StageCFinalHoldoutIdentityIndexError('mapping workbook SHA-256 differs from Task-28 seal')
    headers, rows, worksheet = _read_mapping(mapping_workbook_path); columns=_resolve_columns(headers)
    if len(rows)!=EXPECTED_SAMPLE_COUNT: raise StageCFinalHoldoutIdentityIndexError(f'mapping data-row count mismatch: expected={EXPECTED_SAMPLE_COUNT}, actual={len(rows)}')
    git=_git_provenance(repo_root)
    with zipfile.ZipFile(html_archive_path,'r',allowZip64=True) as archive:
        bad=archive.testzip()
        if bad is not None: raise StageCFinalHoldoutIdentityIndexError(f'HTML archive integrity failure: {bad}')
        infos=[x for x in archive.infolist() if not x.is_dir()]
        if len(infos)!=EXPECTED_SAMPLE_COUNT: raise StageCFinalHoldoutIdentityIndexError(f'HTML member count mismatch: {len(infos)}')
        inventory=[]; exact={}; numeric={}
        for info in sorted(infos,key=lambda x:x.filename):
            if not _safe_member_name(info.filename): raise StageCFinalHoldoutIdentityIndexError(f'unsafe HTML member: {info.filename!r}')
            basename=PurePosixPath(info.filename).name
            if PurePosixPath(basename).suffix.casefold()!='.txt': raise StageCFinalHoldoutIdentityIndexError(f'unexpected HTML member: {info.filename}')
            k=basename.casefold()
            if k in exact: raise StageCFinalHoldoutIdentityIndexError(f'duplicate HTML basename: {basename}')
            exact[k]=info
            stem=PurePosixPath(basename).stem
            try: sk,si=_normalize_serial(stem)
            except StageCFinalHoldoutIdentityIndexError: sk,si=stem.casefold(),None
            if si is not None:
                if sk in numeric: raise StageCFinalHoldoutIdentityIndexError(f'duplicate numeric HTML serial: {sk}')
                numeric[sk]=info
            inventory.append({'name':info.filename,'size_bytes':info.file_size,'compressed_size_bytes':info.compress_size,'crc32':f'{info.CRC:08x}','compression_method':info.compress_type})
        if canonical_hash(inventory)!=EXPECTED_HTML_MEMBER_INVENTORY_SHA256: raise StageCFinalHoldoutIdentityIndexError('HTML member inventory differs from Task-28 seal')
        seen_serials=set(); seen_members=set(); sample_ids=set(); labels=Counter(); records=[]
        for rowno,row in enumerate(rows,start=2):
            get=lambda s: row[columns[s]] if columns[s] < len(row) else None
            serial_key,serial_int=_normalize_serial(get('serial'))
            if serial_key in seen_serials: raise StageCFinalHoldoutIdentityIndexError(f'duplicate serial at row {rowno}: {serial_key}')
            seen_serials.add(serial_key); raw_url=str(get('url')).strip(); norm_url,host=_normalize_url(raw_url); label=_parse_label(get('label')); labels[label]+=1
            info=None
            if 'html' in columns and get('html') is not None and str(get('html')).strip():
                base=PurePosixPath(str(get('html')).replace('\\','/')).name
                if not base.casefold().endswith('.txt'): base += '.txt'
                info=exact.get(base.casefold())
            if info is None: info=numeric.get(serial_key) or exact.get(f'{serial_key}.txt'.casefold())
            if info is None: raise StageCFinalHoldoutIdentityIndexError(f'cannot resolve HTML member for serial {serial_key}')
            if info.filename in seen_members: raise StageCFinalHoldoutIdentityIndexError(f'multiple mapping rows resolve to {info.filename}')
            seen_members.add(info.filename); sid=_sample_id(serial_key,serial_int)
            if sid in sample_ids: raise StageCFinalHoldoutIdentityIndexError(f'duplicate sample_id: {sid}')
            sample_ids.add(sid)
            records.append({'sample_id':sid,'serial_key':serial_key,'serial_number':serial_int,'label':label,
                'url_sha256':hashlib.sha256(raw_url.encode()).hexdigest(),'normalized_url_sha256':hashlib.sha256(norm_url.encode()).hexdigest(),
                'hostname_sha256':hashlib.sha256(host.encode()).hexdigest(),'html_member_name':info.filename,'html_basename':PurePosixPath(info.filename).name,
                'html_size_bytes':info.file_size,'html_crc32':f'{info.CRC:08x}','html_sha256':sha256_zip_member(archive,info)})
    if len(seen_members)!=EXPECTED_SAMPLE_COUNT: raise StageCFinalHoldoutIdentityIndexError('not all sealed HTML captures were mapped')
    if dict(labels)!=EXPECTED_CLASS_COUNTS: raise StageCFinalHoldoutIdentityIndexError(f'class-count mismatch: expected={EXPECTED_CLASS_COUNTS}, actual={dict(labels)}')
    records.sort(key=lambda x:(x['serial_number'] is None,x['serial_number'] or 0,x['serial_key']))
    html_counts=Counter(x['html_sha256'] for x in records); url_counts=Counter(x['normalized_url_sha256'] for x in records)
    index_core={'candidate_id':EXPECTED_CANDIDATE_ID,'source_id':EXPECTED_SOURCE_ID,'doi':EXPECTED_DOI,'sample_count':len(records),'class_counts':{'0':labels[0],'1':labels[1]},'records':records}
    index={'schema_version':INDEX_SCHEMA,'status':'PASS','stage':'C','role':'FINAL_HOLDOUT','protocol_id':'low-fpr-generalization-v1','research_only':True,'deployment_authorized':False,**index_core,
        'raw_urls_persisted':False,'mapping_labels_accessed_for_identity_index':True,'mapping_labels_used_for_model_selection':False,
        'html_capture_bytes_accessed_for_identity_hashing':True,'html_features_extracted':False,'model_scoring_performed':False,'metrics_computed':False,'error_analysis_performed':False,
        'record_set_sha256':canonical_hash(records),'index_identity_sha256':canonical_hash(index_core),
        'identity_bindings':{'task28_local_integrity_seal_sha256':EXPECTED_TASK28_SEAL_SHA256,'task28_local_artifact_set_sha256':EXPECTED_TASK28_ARTIFACT_SET_SHA256,'html_archive_sha256':EXPECTED_HTML_SHA256,'mapping_workbook_sha256':EXPECTED_MAPPING_SHA256,**git},
        'next_gate':'AUDIT_STAGE_C_FINAL_HOLDOUT_CONTAMINATION_AGAINST_DEVELOPMENT_AND_CONSUMED_STAGE_B_TEST'}
    report_core={'candidate_id':EXPECTED_CANDIDATE_ID,'worksheet_name':worksheet,'headers':headers,'normalized_headers':[_normalize_header(x) for x in headers],
        'resolved_columns':{k:{'column_index_zero_based':v,'header':headers[v]} for k,v in sorted(columns.items())},'mapping_data_rows':len(rows),'mapped_html_members':len(seen_members),
        'unique_serials':len(seen_serials),'unique_sample_ids':len(sample_ids),'class_counts':{'legitimate':labels[0],'phishing':labels[1]},'raw_urls_persisted':False,
        'exact_html_duplicate_groups_within_holdout':sum(1 for c in html_counts.values() if c>1),'exact_html_duplicate_samples_within_holdout':sum(c for c in html_counts.values() if c>1),
        'normalized_url_duplicate_groups_within_holdout':sum(1 for c in url_counts.values() if c>1),'normalized_url_duplicate_samples_within_holdout':sum(c for c in url_counts.values() if c>1)}
    report={'schema_version':REPORT_SCHEMA,'status':'PASS','stage':'C','role':'FINAL_HOLDOUT','inspection_mode':'MAPPING_SCHEMA_IDENTITY_ONLY_NO_MODEL_SCORING_NO_FEATURE_EXTRACTION','research_only':True,'deployment_authorized':False,**report_core,
        'mapping_schema_inspected':True,'mapping_labels_accessed':True,'mapping_labels_used_only_for_identity_and_class_reconciliation':True,
        'html_capture_contents_accessed_for_identity_hashing':True,'html_feature_extraction_authorized':False,'model_scoring_authorized':False,'metrics_authorized':False,
        'final_holdout_identity_index_built':True,'record_set_sha256':index['record_set_sha256'],'index_identity_sha256':index['index_identity_sha256'],'report_evidence_sha256':canonical_hash(report_core),'next_gate':index['next_gate']}
    return index, report
